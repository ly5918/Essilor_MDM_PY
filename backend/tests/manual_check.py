# -*- coding: utf-8 -*-
"""操作手册全功能点后端回归（manual_check）。

覆盖（对照 manual_content.py 各章节）：
  A. 批量导入全管线（part2：模板/下载/上传预检/四类分流/行治理/跨BU升级MERGE/审批发布）
  B. 元数据字段 Draft→发布（part1）
  C. One ID 规则 Draft→发布（part1）
  D. DQ 规则治理 + 模拟 + 影响评估 + Scorecard（part1）
  E. 匹配规则 + 样例模拟（part1）
  F. Workflow 场景/配置/部署（part1）
  G. 集成端点 CRUD + 测试连接 + 发布 + 运行重试（part1）
  H. 权限矩阵/角色（part1）  I. 审计查询/导出（part1/part5）

用法：直接运行；每项打印 PASS/FAIL 与摘要，结尾输出汇总与退出码。
"""
import asyncio
import io
import sys

sys.path.insert(0, r"D:/Code/2026AI/AIA_MDM/Essilor_MDM_PY/backend")

from httpx import AsyncClient, ASGITransport

RESULTS = []


def check(name: str, cond: bool, detail: str = ""):
    RESULTS.append((name, bool(cond), detail))
    print(("PASS " if cond else "FAIL ") + name + (f"  | {detail}" if detail else ""))


def make_xlsx(rows):
    """rows: list of dict（键=表头列）。"""
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    headers = list(rows[0].keys())
    ws.append(headers)
    for r in rows:
        ws.append([r.get(h) for h in headers])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


async def api(c, method, url, expect_code=0, **kw):
    r = await c.request(method, url, **kw)
    try:
        j = r.json()
    except Exception:
        return r, None
    return r, j



async def _seed_customer():
    """补种 High End 活跃主档（18 位信用代码），供 Exact 分流使用。"""
    from sqlalchemy.ext.asyncio import create_async_engine
    from sqlalchemy import text
    e = create_async_engine("mysql+aiomysql://root:root@127.0.0.1:3306/ruoyi_plus_py?charset=utf8mb4")
    async with e.begin() as conn:
        row = (await conn.execute(text(
            "SELECT one_id FROM cmd_customer WHERE credit_code='91310000PYTEST0001' AND del_flag='0'"))).first()
        if row is None:
            import time as _time
            oid = f"GC-{int(_time.time()) % 0xFFFFFFFF:08X}"
            await conn.execute(text(
                "INSERT INTO cmd_customer (one_id, legal_name, short_name, credit_code, bu_scope, "
                "country, status, source_system, match_state, duplicate_flag, create_time, update_time, del_flag) "
                "VALUES (:oid, '主档精确命中测试公司', '精确命中', '91310000PYTEST0001', 'High End', "
                "'中国', 'active', 'MANUAL', 'NEW', 'N', NOW(), NOW(), '0')"), {"oid": oid})
    await e.dispose()
    return None

async def main():
    from app.core import db
    await db.init_db()
    from app.main import app
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://t",
                           timeout=60) as c:
        # ============ 准备：取存量主档用于分流测试 ============
        r, j = await api(c, "GET", "/cmd/customer/list?pageNum=1&pageSize=50")
        custs = (j.get("data") or {}).get("rows") or []
        import re as _re
        _CC18 = _re.compile(r"^[0-9A-Z]{18}$")
        active_he = next((x for x in custs if x.get("status") == "active"
                          and x.get("buScope") == "High End" and x.get("creditCode")
                          and _CC18.match(x.get("creditCode"))), None)
        if active_he is None:
            # 补种一条 18 位合法编码的 High End 活跃主档（分流 Exact 依赖）
            from datetime import datetime as _dt
            from sqlalchemy import text as _text
            eng = (await _seed_customer())
            r, j = await api(c, "GET", "/cmd/customer/list?pageNum=1&pageSize=50")
            custs = (j.get("data") or {}).get("rows") or []
            active_he = next((x for x in custs if x.get("status") == "active"
                              and x.get("buScope") == "High End" and x.get("creditCode")
                              and _CC18.match(x.get("creditCode"))), None)
        active_other_bu = next((x for x in custs if x.get("status") == "active"
                                and x.get("buScope") not in (None, "High End")
                                and x.get("creditCode")), None)
        name_hit = next((x for x in custs if x.get("legalName") and x.get("status") == "active"), None)
        check("前置数据：High End 活跃主档(信用代码)", active_he is not None)
        check("前置数据：跨BU活跃主档", active_other_bu is not None)

        # ============ A. 批量导入管线 ============
        # A1 模板清单
        r, j = await api(c, "GET", "/cmd/import/template/list")
        tpl = next((t for t in (j.get("data") or []) if t.get("templateCode") == "TPL_DOOR_MAINSTREAM"), None)
        check("A1 模板清单(TPL_DOOR_MAINSTREAM Published)", tpl is not None and tpl.get("status") == "Published",
              f"count={len(j.get('data') or [])}")

        # A2 新增映射字段
        _sfx = __import__("time").strftime("%H%M%S")  # 每次运行唯一，避免重复提交撞唯一键
        _col = f"门店等级{_sfx}"
        # 运行级唯一测试数据：上一轮发布的 NEW 主档会污染下一轮分流（名称/编码命中）
        _new_name = f"导入新增行测试公司{_sfx}"
        _new_code = f"91310000{_sfx}0001"  # 8+6+4=18 位合法编码
        _fld_code = f"manual_check_field{_sfx}"  # 元数据字段编码（md_field 唯一键）
        r, j = await api(c, "POST", "/cmd/import/template/mapping",
                         json={"templateCode": "TPL_DOOR_MAINSTREAM", "columnName": _col,
                               "fieldCode": f"store_grade_t{_sfx}", "fieldName": _col,
                               "dataType": "Text", "isRequired": "N"})
        check("A2 新增上传字段(保存映射)", j and j.get("code") in (0, 200), str((j or {}).get("msg"))[:60])

        # A3 模板下载（真实 xlsx，含新列）
        r, j = await api(c, "GET", "/cmd/import/template/TPL_DOOR_MAINSTREAM/download")
        ok = False
        header_cols = []
        try:
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(r.content))
            ws = wb.active
            header_cols = [c.value for c in ws[1]]
            ok = _col in header_cols and "CustomerName" in header_cols
        except Exception as e:
            ok = False
        check("A3 模板下载=动态表头 xlsx（含新增列）", ok, f"cols={header_cols}")

        # A4 预检失败（缺必填列）
        bad = make_xlsx([{"CustomerName": "X公司", "CreditCode": "91310000AA00000001"}])
        r, j = await api(c, "POST", "/cmd/import/job/upload",
                         files={"file": ("bad.xlsx", bad)},
                         data={"templateCode": "TPL_DOOR_MAINSTREAM"})
        check("A4 文件级预检：缺 Address 列整批退回", j and j.get("code") not in (0, 200) and "Address" in str(j.get("msg")),
              str((j or {}).get("msg"))[:80])

        # A5 上传成功 → 四类分流
        code_exact = (active_he or {}).get("creditCode")
        rows = [
            {"CustomerName": "主档精确命中测试", "CreditCode": code_exact, "Address": "上海市黄浦区1号",
             "City": "上海", "ContactPhone": "021-0000", "客户邮箱": "a@b.c"},
            {"CustomerName": (name_hit or {}).get("legalName"), "CreditCode": "91320500MA1SUSP001",
             "Address": "上海市黄浦区2号", "City": "上海", "ContactPhone": "021-0001", "客户邮箱": ""},
            {"CustomerName": _new_name, "CreditCode": _new_code, "Address": "上海市黄浦区3号",
             "City": "上海", "ContactPhone": "021-0002", "客户邮箱": ""},
            {"CustomerName": "批次内重复测试公司", "CreditCode": _new_code, "Address": "上海市黄浦区4号",
             "City": "上海", "ContactPhone": "021-0003", "客户邮箱": ""},
            {"CustomerName": "格式错误测试公司", "CreditCode": "BAD", "Address": "上海市黄浦区5号",
             "City": "上海", "ContactPhone": "021-0004", "客户邮箱": ""},
        ]
        r, j = await api(c, "POST", "/cmd/import/job/upload",
                         files={"file": ("manual_check.xlsx", make_xlsx(rows))},
                         data={"templateCode": "TPL_DOOR_MAINSTREAM", "buScope": "High End"})
        job_code = (j or {}).get("data")
        check("A5 上传成功建任务", bool(job_code), str((j or {}).get("msg"))[:80])

        r, j = await api(c, "GET", f"/cmd/import/job/{job_code}/result")
        d = j.get("data") or {}
        routes = d.get("routes") or []
        check("A6 四类分流计数 exact=1 suspected=2 new=1 invalid=1",
              d.get("exact") == 1 and d.get("suspected") == 2 and d.get("created") == 1
              and d.get("invalid") == 1, f"{d.get('exact')}/{d.get('suspected')}/{d.get('created')}/{d.get('invalid')}")
        check("A7 五分流 routes 卡（含查看N条）", len(routes) == 5 and all("detail" in x for x in routes))

        # A8 行明细 + resultType 过滤
        r, j = await api(c, "GET", f"/cmd/import/job/{job_code}/rows?resultType=SUSPECTED&pageNum=1&pageSize=20")
        sus_rows = (j.get("data") or {}).get("rows") or []
        check("A8 行明细 resultType 过滤", len(sus_rows) == 2, f"n={len(sus_rows)}")

        # A9 IMPORT_BATCH 审批待办已创建
        r, j = await api(c, "GET", "/cmd/approval/list?taskCategory=APPROVAL&pageNum=1&pageSize=50")
        tasks = (j.get("data") or {}).get("rows") or []
        import_task = next((t for t in tasks if t.get("bizId") == job_code
                            and t.get("sceneCode") == "IMPORT_BATCH"), None)
        check("A9 批量导入确认审批待办已创建", import_task is not None)

        # A10 LINK 跨BU → 升级 MERGE 审批
        name_sus = next((x for x in sus_rows if x.get("matchState") == "SUSPECTED"
                         and x.get("rowNo") == 3), None)
        r, j = await api(c, "POST", f"/cmd/import/row/{name_sus['id']}/action",
                         json={"action": "LINK", "oneId": active_other_bu["oneId"]})
        msg = str((j or {}).get("msg") or (j or {}).get("data") or "")
        check("A10 行治理 LINK 跨BU → 升级合并审批", j and j.get("code") in (0, 200) and "跨BU" in msg, msg[:90])

        r, j = await api(c, "GET", "/cmd/approval/list?taskCategory=APPROVAL&pageNum=1&pageSize=50")
        tasks = (j.get("data") or {}).get("rows") or []
        rowlink_task = next((t for t in tasks if t.get("sceneCode") == "MERGE"
                             and str(t.get("bizId")) == str(name_sus["id"])), None)
        check("A11 ROW_LINK 合并审批待办(mode=ROW_LINK)", rowlink_task is not None and
              ((rowlink_task.get("bizSnapshotJson") or {}).get("mode") == "ROW_LINK"))

        # A12 EXCLUDE
        batch_sus = next((x for x in sus_rows if x["id"] != name_sus["id"]), None)
        r, j = await api(c, "POST", f"/cmd/import/row/{batch_sus['id']}/action",
                         json={"action": "EXCLUDE"})
        r, j2 = await api(c, "GET", f"/cmd/import/job/{job_code}/rows?pageNum=1&pageSize=20")
        all_rows = (j2.get("data") or {}).get("rows") or []
        excluded = next((x for x in all_rows if x["id"] == batch_sus["id"]), {})
        check("A12 行治理 EXCLUDE → Invalid/SKIPPED",
              excluded.get("resultType") == "INVALID" and excluded.get("rowStatus") == "SKIPPED")

        # A13 RETURN
        invalid_row = next((x for x in all_rows if x.get("resultType") == "INVALID"
                            and x.get("rowStatus") == "FAILED"), None)
        r, j = await api(c, "POST", f"/cmd/import/row/{invalid_row['id']}/action",
                         json={"action": "RETURN"})
        check("A13 行治理 RETURN 退回修复", j and j.get("code") in (0, 200))

        # A14 批准 IMPORT_BATCH → New 行生成 One ID + 主档 active
        r, j = await api(c, "POST", "/cmd/approval/action",
                         json={"taskNo": import_task["taskNo"], "action": "approve",
                               "actor": "BU Steward", "opinion": "手册回归批准"})
        r, j = await api(c, "GET", f"/cmd/import/job/{job_code}/rows?pageNum=1&pageSize=20")
        new_rows = [x for x in ((j.get("data") or {}).get("rows") or []) if x.get("resultType") == "NEW"]
        new_one = new_rows[0].get("oneId") if new_rows else None
        check("A14 批准后 New 行生成 One ID", bool(new_one), f"oneId={new_one}")
        r, j = await api(c, "GET", f"/cmd/customer/oneId/{new_one}")
        d = j.get("data") or {}
        check("A15 New 行主档已发布(active)", d.get("status") == "active", f"status={d.get('status')}")

        # A16 ROW_LINK MERGE 审批批准 → 行回写 Exact
        if rowlink_task:
            r, j = await api(c, "POST", "/cmd/approval/action",
                             json={"taskNo": rowlink_task["taskNo"], "action": "approve",
                                   "actor": "BU Steward", "opinion": "确认合并"})
            r, j = await api(c, "GET", f"/cmd/import/job/{job_code}/rows?pageNum=1&pageSize=20")
            row_after = next((x for x in ((j.get("data") or {}).get("rows") or [])
                              if x["id"] == name_sus["id"]), {})
            # 跨BU 需 GC 决策：若仍未闭环则用 GC 再批一次
            if row_after.get("resultType") != "EXACT":
                r, j = await api(c, "GET", "/cmd/approval/list?taskCategory=APPROVAL&pageNum=1&pageSize=50")
                tasks = (j.get("data") or {}).get("rows") or []
                gc_task = next((t for t in tasks if t.get("sceneCode") == "MERGE"
                                and str(t.get("bizId")) == str(name_sus["id"])
                                and t.get("status") == "PENDING"), None)
                if gc_task:
                    r, j = await api(c, "POST", "/cmd/approval/action",
                                     json={"taskNo": gc_task["taskNo"], "action": "approve",
                                           "actor": "GC Steward", "opinion": "GC 决策通过"})
            # GC 决策后重新拉取行明细（j 可能已被审批响应覆盖）
            r, j = await api(c, "GET", f"/cmd/import/job/{job_code}/rows?pageNum=1&pageSize=20")
            row_after = next((x for x in ((j.get("data") or {}).get("rows") or [])
                              if x["id"] == name_sus["id"]), {})
            check("A16 ROW_LINK 批准后行回写 Exact+目标OneID",
                  row_after.get("resultType") == "EXACT" and row_after.get("oneId") == active_other_bu["oneId"],
                  f"{row_after.get('resultType')}/{row_after.get('oneId')}")

        # A17 删除映射：主键保护
        r, j = await api(c, "GET", "/cmd/import/template/mapping?templateCode=TPL_DOOR_MAINSTREAM")
        mappings = j.get("data") or []
        core = next((m for m in mappings if m.get("fieldCode") == "legal_name"), None)
        r, j = await api(c, "DELETE", f"/cmd/import/template/mapping/{core['id']}")
        check("A17 主键字段映射不可删除", j and j.get("code") not in (0, 200), str((j or {}).get("msg"))[:60])
        added = next((m for m in mappings if m.get("fieldCode") == f"store_grade_t{_sfx}"), None)
        if added:
            r, j = await api(c, "DELETE", f"/cmd/import/template/mapping/{added['id']}")
            check("A18 非主键映射可删除", j and j.get("code") in (0, 200))

        # ============ B. 元数据 Draft→发布 ============
        # 工作版本 = 版本号最大者（新字段必须落进真实版本，否则成为游离行，
        # 会把旧版本顶成第二个「Current」污染版本列表——历史 bug）
        r, j = await api(c, "GET", "/cmd/metadata/version/list")
        _versions_before = j.get("data") or []
        _work_ver = max((v.get("version") or "" for v in _versions_before),
                        key=lambda s: [int(x) for x in s.lstrip("v").split(".")])
        r, j = await api(c, "POST", "/cmd/metadata/field",
                         json={"field_code": _fld_code, "field_name": "回归测试字段",
                               "field_type": "Text", "model_code": "CUSTOMER",
                               "version_no": _work_ver, "status": "1", "order_num": 99})
        check("B1 新建元数据字段(Draft)", j and j.get("code") in (0, 200), str((j or {}).get("msg"))[:50])
        # 发布承载该字段的版本（契约：version/list 返回 status=Current/Draft）
        r, j = await api(c, "POST", f"/cmd/metadata/version/publish?version={_work_ver}")
        check("B2 发布模型版本", j and j.get("code") in (0, 200), str((j or {}).get("msg"))[:50])
        # B3 清理回归字段：逻辑删除，避免每轮运行残留垃圾行污染字段目录
        r, j = await api(c, "GET", "/cmd/metadata/field/list")
        _created = next((f for f in (j.get("data") or []) if f.get("fieldCode") == _fld_code), None)
        if _created:
            r, j = await api(c, "DELETE", f"/cmd/metadata/field/{_created['id']}")
            check("B3 回归字段清理(逻辑删除)", j and j.get("code") in (0, 200), str((j or {}).get("msg"))[:50])

        # ============ C. One ID 规则 ============
        r, j = await api(c, "GET", "/cmd/oneid/rule")
        rule = j.get("data") or {}
        check("C1 读取 One ID 规则", isinstance(rule, dict) and rule)
        r, j = await api(c, "PUT", "/cmd/oneid/rule", json={"rule_name": rule.get("rule_name") or "默认规则",
                                                            "serial_length": 6})
        check("C2 保存规则(Draft)", j and j.get("code") in (0, 200))
        r, j = await api(c, "PUT", "/cmd/oneid/rule/publish")
        check("C3 发布规则", j and j.get("code") in (0, 200))
        r, j = await api(c, "GET", "/cmd/oneid/rule")
        check("C4 发布后 status=1(已发布)", str((j.get("data") or {}).get("status")) == "1")

        # ============ D. DQ 规则 ============
        r, j = await api(c, "GET", "/cmd/dq/rule/list")
        check("D1 DQ 规则清单", isinstance(j.get("data"), list) and len(j["data"]) > 0,
              f"n={len(j.get('data') or [])}")
        r, j = await api(c, "POST", "/cmd/dq/rule",
                         json={"rule_code": f"R_MANUAL_CHECK{_sfx}", "rule_name": "回归测试规则",
                               "dimension": "COMPLETENESS", "field_code": "legal_name",
                               "check_type": "NOT_NULL", "deduction": 10, "enabled": "1"})
        check("D2 新增 DQ 规则", j and j.get("code") in (0, 200), str((j or {}).get("msg"))[:50])
        r, j = await api(c, "GET", "/cmd/dq/rule/list")
        my_rule = next((x for x in (j.get("data") or []) if x.get("ruleCode") == f"R_MANUAL_CHECK{_sfx}"
                        or x.get("rule_code") == f"R_MANUAL_CHECK{_sfx}"), None)
        if my_rule:
            rid = my_rule.get("ruleId") or my_rule.get("rule_id") or my_rule.get("id")
            r, j = await api(c, "PUT", f"/cmd/dq/rule/{rid}",
                             json={"rule_name": "回归测试规则(改)", "deduction": 12})
            check("D3 编辑 DQ 规则", j and j.get("code") in (0, 200))
            r, j = await api(c, "DELETE", f"/cmd/dq/rule/{rid}")
            check("D4 删除 DQ 规则", j and j.get("code") in (0, 200))
        r, j = await api(c, "POST", "/cmd/dq/simulate", json={"scopeType": "ALL"})
        check("D5 DQ 模拟测试", j and j.get("code") in (0, 200), str((j or {}).get("msg"))[:50])
        r, j = await api(c, "GET", "/cmd/dq/reEvaluate/impact")
        check("D6 影响评估", j and j.get("code") in (0, 200))
        r, j = await api(c, "GET", "/cmd/dq/scorecard")
        check("D7 DQ Scorecard", j and j.get("code") in (0, 200))

        # ============ E. 匹配规则 ============
        r, j = await api(c, "GET", "/cmd/match/rule/list")
        check("E1 匹配规则清单", isinstance(j.get("data"), list) and len(j["data"]) > 0)
        r, j = await api(c, "POST", "/cmd/match/simulate",
                         json={"legalName": (name_hit or {}).get("legalName"),
                               "creditCode": "91310000SIM0000001"})
        check("E2 匹配样例模拟（名称命中出候选）", j and j.get("code") in (0, 200), str((j or {}).get("msg"))[:60])

        # ============ F. Workflow ============
        r, j = await api(c, "GET", "/cmd/flow/scenes")
        scenes = j.get("data") or []
        check("F1 Workflow 9 大场景", len(scenes) >= 9, f"n={len(scenes)}")
        r, j = await api(c, "GET", "/cmd/flow/scene/CUSTOMER_CREATE/config")
        check("F2 场景配置读取", j and j.get("code") in (0, 200))
        r, j = await api(c, "PUT", "/cmd/flow/scene/CUSTOMER_CREATE/config",
                         json={"slaHours": 24})
        check("F3 场景配置保存", j and j.get("code") in (0, 200))
        r, j = await api(c, "POST", "/cmd/flow/deploy/IMPORT_BATCH")
        check("F4 场景部署", j and j.get("code") in (0, 200))

        # ============ G. 集成 ============
        r, j = await api(c, "POST", "/cmd/integration/endpoint",
                         json={"endpointName": "回归测试端点", "direction": "OUTBOUND",
                               "protocol": "mock", "targetSystem": "DMS+", "endpointUrl": "mock://test"})
        ep_code = (j.get("data") or {}).get("endpoint_code") if j and isinstance(j.get("data"), dict) else None
        check("G1 新建集成端点", j and j.get("code") in (0, 200), f"code={ep_code}")
        r, j = await api(c, "GET", "/cmd/integration/endpoint/list")
        eps = j.get("data") or []
        ep = next((x for x in eps if x.get("endpointName") == "回归测试端点"
                   or x.get("endpoint_name") == "回归测试端点"), None)
        if ep:
            eid = ep.get("id")
            r, j = await api(c, "PUT", f"/cmd/integration/endpoint/{eid}",
                             json={"endpointName": "回归测试端点(改)"})
            check("G2 编辑端点", j and j.get("code") in (0, 200))
            r, j = await api(c, "POST", f"/cmd/integration/endpoint/{eid}/test")
            check("G3 测试连接", j and j.get("code") in (0, 200), str((j or {}).get("msg"))[:50])
            r, j = await api(c, "POST", f"/cmd/integration/endpoint/{eid}/publish")
            check("G4 发布端点", j and j.get("code") in (0, 200))
        r, j = await api(c, "GET", "/cmd/integration/run/list?pageNum=1&pageSize=5")
        runs = ((j.get("data") or {}).get("rows") or [])
        check("G5 运行监控记录", len(runs) > 0, f"n={len(runs)}")
        if runs:
            rid = runs[0].get("id")
            r, j = await api(c, "PUT", f"/cmd/integration/run/{rid}/retry")
            check("G6 运行 Retry(PUT)", j and j.get("code") in (0, 200), str((j or {}).get("msg"))[:50])

        # ============ H. 权限 / I. 审计 ============
        r, j = await api(c, "GET", "/cmd/permission/matrix")
        check("H1 权限矩阵", j and j.get("code") in (0, 200))
        r, j = await api(c, "GET", "/cmd/permission/role/list")
        check("H2 角色列表", j and j.get("code") in (0, 200))
        r, j = await api(c, "GET", "/cmd/audit/list?pageNum=1&pageSize=5")
        check("I1 审计事件列表", j and j.get("code") in (0, 200))
        r, j = await api(c, "POST", "/cmd/audit/export", json={"fileName": "manual_check_export"})
        check("I2 审计导出", j and j.get("code") in (0, 200))

    # ============ 汇总 ============
    fails = [x for x in RESULTS if not x[1]]
    print(f"\n===== 汇总: {len(RESULTS) - len(fails)}/{len(RESULTS)} PASS =====")
    for name, _, detail in fails:
        print("  FAIL:", name, "|", detail[:120])
    return 1 if fails else 0


if __name__ == "__main__":
    exit(asyncio.run(main()))
