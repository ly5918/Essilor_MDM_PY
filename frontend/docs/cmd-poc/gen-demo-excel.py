# -*- coding: utf-8 -*-
"""
生成批量导入演示数据（5 行覆盖四类分流）。

为什么需要"每次换新的客户名称 + 信用代码"：
    New 行在审批通过后会被发布进客户主档（名称与信用代码都会入库）。
    如果演示 Excel 一直用同一套名称/代码，第二次上传时这行就会命中存量主档
    → 变成 Exact 或 Suspected，不再产生 New，也就不会触发「批量导入确认」审批。
    所以这里每次挑一个主档中不存在的名称，并随机生成一个未使用过的信用代码。

输出：docs/cmd-poc/批量导入_演示数据.xlsx
     （若该文件正被 Excel 打开而无法写入，自动改存为 批量导入_演示数据_新.xlsx）
用法：python gen-demo-excel.py
"""
import random
import string
import subprocess
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

HERE = Path(__file__).resolve().parent
OUT = Path(sys.argv[1]) if len(sys.argv) > 1 else HERE / "批量导入_演示数据.xlsx"

MYSQL = r"D:\Code\mysql-8.0.37-winx64\bin\mysql.exe"

HEADERS = ["CustomerName", "CreditCode", "Address", "City", "ContactPhone"]

ALPHABET = string.digits + "ABCDEFGHIJKLMNOPQRSTUVWXYZ"

# 候选新客户名称池（按顺序取第一个主档中不存在的）
NAME_POOL = [
    "无锡明亮视光有限公司",
    "常州汇视眼镜有限公司",
    "南通晶明光学有限公司",
    "嘉兴博视眼镜有限公司",
    "温州瓯江光学有限公司",
    "金华双龙视光有限公司",
    "绍兴大光明眼镜有限公司",
    "台州椒江眼镜有限公司",
]


def mysql_scalar(sql: str) -> list:
    """查询主档，失败时返回空列表（脚本仍可离线生成）。"""
    try:
        res = subprocess.run(
            [MYSQL, "-uroot", "-proot", "-D", "ruoyi_plus", "-N", "-e", sql],
            capture_output=True, text=True, timeout=20,
        )
        return [x.strip() for x in (res.stdout or "").splitlines() if x.strip()]
    except Exception as e:  # noqa: BLE001
        print(f"[gen-demo-excel] 读取主档失败（忽略）：{e}")
        return []


def pick_new_customer(used_names: set, used_codes: set) -> tuple:
    name = next((n for n in NAME_POOL if n not in used_names), None)
    if name is None:  # 池子用尽，追加编号保证唯一
        name = f"{NAME_POOL[0][:-4]}{random.randint(1000, 9999)}有限公司"
    for _ in range(50):
        code = "91330200MA2" + "".join(random.choice(ALPHABET) for _ in range(7))
        if code not in used_codes:
            return name, code
    return name, "91330200MA2" + "".join(random.choice(ALPHABET) for _ in range(7))


def main() -> None:
    used_names = set(mysql_scalar("SELECT legal_name FROM cmd_customer;"))
    used_codes = set(mysql_scalar("SELECT credit_code FROM cmd_customer WHERE credit_code IS NOT NULL;"))
    new_name, new_code = pick_new_customer(used_names, used_codes)

    rows = [
        # 1) 命中存量主档信用代码 → Exact
        ["上海清视眼镜有限公司", "91310000MA1K35XX8X", "上海市静安区南京西路100号", "上海", "021-60000001"],
        # 2) 主档不存在 → New（审批通过后生成 One ID）
        [new_name, new_code, "江苏省无锡市梁溪区人民中路800号", "无锡", "0510-60000008"],
        # 3) 名称与主档相同、信用代码不同 → Suspected
        ["苏州新视野眼镜有限公司", "91320500MA1SUSP009", "苏州市工业园区星湖街300号", "苏州", "0512-60000003"],
        # 4) 与本批第 2 行信用代码重复 → Suspected（批次内去重）
        [new_name, new_code, "江苏省无锡市梁溪区人民中路800号", "无锡", "0510-60000008"],
        # 5) 必填地址缺失 + 信用代码非 18 位 → Invalid
        ["残缺点客户", "9132BAD", None, None, None],
    ]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Template"

    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="176C9F")
    for col, title in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=col, value=title)
        cell.font = head_font
        cell.fill = head_fill
        cell.alignment = Alignment(horizontal="center")

    for r, row in enumerate(rows, start=2):
        for c, value in enumerate(row, start=1):
            ws.cell(row=r, column=c, value=value)

    widths = [30, 24, 34, 10, 18]
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(ord("A") + i)].width = w

    target = OUT
    try:
        wb.save(target)
    except PermissionError:
        target = OUT.with_name(OUT.stem + "_新" + OUT.suffix)
        wb.save(target)
        print(f"[gen-demo-excel] 原文件被占用（Excel 打开中），已改存为 {target}")

    print(f"[gen-demo-excel] 输出={target}")
    print(f"[gen-demo-excel] 本次 New 行 = {new_name} / {new_code}")


if __name__ == "__main__":
    main()
